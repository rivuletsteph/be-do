# -*- coding: utf-8 -*-
"""The weekly export's pieces: her words out of a row, a clock that reads the
same on every machine, and markdown to .docx with or without pandoc.

Her words only. A row's details are her words, the divider, then the [be•do]
block. The export keeps what is above the divider and nothing of the block —
except blockquoted screen content, which is a transcription, not be•do's voice.
A row that OPENS with the block has no words of hers, so it exports nothing of
it. The legacy --- divider is read as a divider too: 190 W40 rows used it, and
the export leaked their blocks.
"""
import re, shutil, subprocess

from bedo_reader import FACTS

_DIV = re.compile(r'^\s*(?:' + '|'.join(re.escape(d) for d in [FACTS['divider']] + FACTS['legacy_dividers'])
                  + r')\s*$', re.M)
BLOCK = FACTS['be_do_block']


def hers(details):
    if not details:
        return ''
    parts = _DIV.split(details)
    keep = parts[0].strip()
    if keep.startswith(BLOCK):              # no words of hers above the block
        parts, keep = [''] + parts, ''
    elif BLOCK in keep:                      # a block with no divider before it
        keep = keep.split(BLOCK)[0].strip()
    for p in parts[1:]:                      # transcribed screen content survives
        q = [l for l in p.splitlines() if l.startswith('>')]
        if q:
            keep += '\n\n' + '\n'.join(q)
    return keep.strip()


def clock(t):
    """7:05 AM on every platform. %-I is glibc-only and Windows refuses it."""
    return f'{t.hour % 12 or 12}:{t:%M} {"AM" if t.hour < 12 else "PM"}'


def _docx_by_hand(md, out):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.styles['Normal'].font.name = 'Calibri'
    d.styles['Normal'].font.size = Pt(11)

    def runs(p, text):
        for i, part in enumerate(re.split(r'\*\*', text)):
            if part:
                p.add_run(part).bold = (i % 2 == 1)

    para = []

    def flush():
        if para:
            if all(l.startswith('>') for l in para):
                runs(d.add_paragraph(style='Quote'), '\n'.join(l[1:].lstrip() for l in para))
            else:
                runs(d.add_paragraph(), '\n'.join(para))
            para.clear()

    for l in md.split('\n'):
        if l.startswith('# '):
            flush(); d.add_heading(l[2:], 0)
        elif l.startswith('## '):
            flush(); d.add_heading(l[3:], 1)
        elif not l.strip():
            flush()
        else:
            para.append(l)
    flush()
    d.save(out)


def write_docx(md_path, out):
    """pandoc where it is installed, python-docx where it is not. Returns which
    one ran; stops with a plain message when neither is available."""
    if shutil.which('pandoc'):
        subprocess.run(['pandoc', md_path, '-o', out], check=True)
        return 'pandoc'
    try:
        with open(md_path, encoding='utf-8') as fh:
            _docx_by_hand(fh.read(), out)
        return 'python-docx'
    except ImportError:
        raise SystemExit('ABORT: no pandoc and no python-docx — pip install python-docx, '
                         f'or open {md_path} as it is')
